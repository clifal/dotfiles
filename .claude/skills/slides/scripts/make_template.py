#!/usr/bin/env python3
"""assets/template.pptx を作る。版面と色を変えるときは、ここを直して作り直す。

  python3 -I ~/.claude/skills/slides/scripts/make_template.py

python-pptx の白紙のテンプレートから、レイアウトを表紙・章扉・本文・白紙の 4 つにしぼって組む。
書体は BIZ UDGothic、色はカラーユニバーサルデザイン推奨配色セット（ver.4）から選ぶ（../design.md）。
値を変えたら house.skin.json と house.py の定数もそろえる。
"""
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "assets" / "template.pptx"
EMU = 914400
W, H = 10.0, 5.625
FONT = "BIZ UDGothic"
NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
      'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"')
# テーマの色。dk1 は本文の文字、accent1〜3 は区分の色、accent4 は章扉の番号、accent5 は ✕ の赤
COLORS = dict(dk1="333333", lt1="FFFFFF", dk2="1F2D3D", lt2="F2F4F7", accent1="005AFF", accent2="03AF7A",
              accent3="F6AA00", accent4="4DC4FF", accent5="FF4B00", accent6="990099", hlink="005AFF", folHlink="990099")
NAVY = "accent1|lumMod=50000"   # 章扉の地
SUB = "bg1|lumMod=50000"        # フッターとページ番号


def e(v):
    return str(int(round(v * EMU)))


def clr(spec):
    base, *mods = spec.split("|")
    tag = "srgbClr" if re.fullmatch(r"[0-9A-Fa-f]{6}", base) else "schemeClr"
    inner = "".join(f'<a:{k} val="{v}"/>' for k, v in (m.split("=") for m in mods))
    return f'<a:{tag} val="{base}">{inner}</a:{tag}>'


def xml(s):
    return etree.fromstring(s.replace("<p:sp>", f"<p:sp {NS}>", 1).replace("<p:bg>", f"<p:bg {NS}>", 1))


def xfrm(x, y, w, h):
    return f'<a:xfrm><a:off x="{e(x)}" y="{e(y)}"/><a:ext cx="{e(w)}" cy="{e(h)}"/></a:xfrm>'


def ph(sid, name, kind, x, y, w, h, sz, anchor="t", b=False, c="tx1", algn="l", text="", idx=None):
    """レイアウトのプレースホルダー。余白は 0、自動調整なし、行頭記号なし。"""
    ix = f' idx="{idx}"' if idx is not None else ""
    bold = ' b="1"' if b else ""
    if kind == "sldNum":
        run = ('<a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum">'
               '<a:rPr lang="en-US"/><a:t>‹#›</a:t></a:fld>')
    else:
        run = f'<a:r><a:rPr lang="ja-JP" altLang="en-US"/><a:t>{text}</a:t></a:r>'
    return xml(f'''<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr><a:spLocks noGrp="1"/></p:cNvSpPr>
<p:nvPr><p:ph type="{kind}"{ix}/></p:nvPr></p:nvSpPr><p:spPr>{xfrm(x, y, w, h)}</p:spPr>
<p:txBody><a:bodyPr lIns="0" tIns="0" rIns="0" bIns="0" anchor="{anchor}"><a:noAutofit/></a:bodyPr>
<a:lstStyle><a:lvl1pPr marL="0" indent="0" algn="{algn}"><a:buNone/><a:defRPr sz="{int(sz * 100)}"{bold}>
<a:solidFill>{clr(c)}</a:solidFill></a:defRPr></a:lvl1pPr></a:lstStyle><a:p>{run}</a:p></p:txBody></p:sp>''')


def rect(sid, name, x, y, w, h, c):
    """レイアウトに置く飾りの帯。"""
    return xml(f'''<p:sp><p:nvSpPr><p:cNvPr id="{sid}" name="{name}"/><p:cNvSpPr/><p:nvPr userDrawn="1"/></p:nvSpPr>
<p:spPr>{xfrm(x, y, w, h)}<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:solidFill>{clr(c)}</a:solidFill>
<a:ln><a:noFill/></a:ln></p:spPr></p:sp>''')


def theme(prs):
    part = prs.slide_master.part.part_related_by(RT.THEME)
    s = part.blob.decode("utf8")
    scheme = "".join(f'<a:{k}><a:srgbClr val="{v}"/></a:{k}>' for k, v in COLORS.items())
    s = re.sub(r"<a:clrScheme .*?</a:clrScheme>", f'<a:clrScheme name="UD">{scheme}</a:clrScheme>', s, flags=re.S)
    face = (f'<a:latin typeface="{FONT}"/><a:ea typeface="{FONT}"/><a:cs typeface=""/>'
            f'<a:font script="Jpan" typeface="{FONT}"/>')
    s = re.sub(r"<a:fontScheme .*?</a:fontScheme>",
               f'<a:fontScheme name="UD"><a:majorFont>{face}</a:majorFont><a:minorFont>{face}</a:minorFont></a:fontScheme>',
               s, flags=re.S)
    part._blob = s.encode("utf8")


def master(prs):
    m = prs.slide_master
    boxes = {"TITLE": (0.4, 0.18, 9.2, 0.5), "BODY": (0.4, 0.85, 9.2, 4.15), "DATE": (0.4, 5.3, 2.0, 0.22),
             "FOOTER": (2.6, 5.3, 4.0, 0.22), "SLIDE_NUMBER": (8.6, 5.3, 1.0, 0.22)}
    for p in m.placeholders:
        x, y, w, h = boxes[p.placeholder_format.type.name]
        p.left, p.top, p.width, p.height = (int(round(v * EMU)) for v in (x, y, w, h))
        bp = p._element.txBody.find(qn("a:bodyPr"))
        for k in ("lIns", "tIns", "rIns", "bIns"):
            bp.set(k, "0")
        if p.placeholder_format.type.name == "TITLE":
            bp.set("anchor", "b")
    styles = m._element.find(qn("p:txStyles"))
    t = styles.find(qn("p:titleStyle")).find(qn("a:lvl1pPr"))
    t.set("algn", "l")
    t.find(qn("a:defRPr")).set("sz", "1600")
    t.find(qn("a:defRPr")).set("b", "1")
    body = styles.find(qn("p:bodyStyle"))
    for i, lvl in enumerate(body):
        lvl.set("marL", str(int(0.2 * EMU * (i + 1))))
        lvl.set("indent", str(-int(0.2 * EMU)))
        lvl.find(qn("a:defRPr")).set("sz", str([1200, 1100][i] if i < 2 else 1000))


def layouts(prs):
    """白紙のテンプレートのレイアウトから 4 つを残し、中身を組み直す。"""
    keep = {"Title Slide": "表紙", "Section Header": "章扉", "Title Only": "本文", "Blank": "白紙"}
    for lay in list(prs.slide_layouts):
        if lay.name not in keep:
            prs.slide_layouts.remove(lay)
    def num(sid):
        return ph(sid, "ページ番号", "sldNum", 8.6, 5.3, 1.0, 0.22, 8, anchor="ctr", c=SUB, algn="r", idx=12)

    for lay in prs.slide_layouts:
        name = keep[lay.name]
        lay._element.cSld.set("name", name)
        tree = lay._element.cSld.spTree
        for sp in list(tree)[2:]:
            tree.remove(sp)
        if name == "表紙":
            shapes = [rect(2, "左の帯", 0, 0, 0.2, H, "accent1"),
                      rect(3, "右の面", 8.6, 0, W - 8.6, H, "accent1|lumMod=20000|lumOff=80000"),
                      ph(4, "題", "ctrTitle", 0.7, 1.45, 7.5, 1.5, 26, anchor="b", b=True, text="資料の題"),
                      rect(5, "題の下線", 0.7, 3.1, 0.9, 0.05, "accent1"),
                      ph(6, "日付と作成者", "subTitle", 0.7, 3.3, 7.5, 0.8, 12, c="bg1|lumMod=35000", idx=1,
                         text="日付と作成者")]
        elif name == "章扉":
            bg = xml(f"<p:bg><p:bgPr><a:solidFill>{clr(NAVY)}</a:solidFill><a:effectLst/></p:bgPr></p:bg>")
            lay._element.cSld.insert(0, bg)
            shapes = [ph(2, "章題", "title", 0.7, 1.95, 8.6, 0.95, 28, anchor="b", b=True, c="bg1", text="章の題"),
                      rect(3, "章題の下線", 0.7, 3.05, 0.9, 0.05, "accent4"),
                      ph(4, "章の説明", "body", 0.7, 3.25, 8.6, 0.7, 12, c="bg1", idx=1, text="この章でわかること")]
        elif name == "本文":
            shapes = [ph(2, "タイトル", "title", 0.4, 0.18, 9.2, 0.5, 16, anchor="b", b=True, text="タイトル"),
                      rect(3, "タイトルの罫", 0.4, 0.765, 9.2, 0.01, "bg1|lumMod=85000"),
                      rect(4, "タイトルの印", 0.4, 0.75, 0.6, 0.04, "accent1"), num(5)]
        else:
            shapes = [num(2)]
        for sp in shapes:
            tree.append(sp)


def samples(prs):
    """テンプレートを開いた人が見本を見られるよう、表紙・章扉・本文を 1 枚ずつ置く（Builder は外して使う）。"""
    by = {lay.name: lay for lay in prs.slide_layouts}
    texts = {"表紙": ["資料の題（2 行まで）", "2026/01/01\n作成者"], "章扉": ["章の題", "この章でわかること"],
             "本文": ["ページの結論を 1 文で書くタイトル"]}
    for name, vals in texts.items():
        s = prs.slides.add_slide(by[name])
        for p, v in zip(sorted(s.placeholders, key=lambda p: p.placeholder_format.idx), vals):
            p.text_frame.text = v


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = int(W * EMU), int(H * EMU)
    theme(prs)
    master(prs)
    layouts(prs)
    samples(prs)
    cp = prs.core_properties
    cp.author = cp.last_modified_by = cp.title = cp.subject = cp.keywords = cp.comments = cp.category = ""
    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print("wrote", OUT)


if __name__ == "__main__":
    main()
