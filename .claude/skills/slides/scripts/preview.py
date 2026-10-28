#!/usr/bin/env python3
"""Keynote が使えないときの代わりの目視。house.py が作る図形（矩形・楕円・三角・矢羽・線・文字）を HTML に描き直し、
3 ページずつの画像にする。

  python3 -I ~/.claude/skills/slides/scripts/preview.py <pptx> <出力先フォルダ>

書体と行の折り返しは PowerPoint・Keynote と少し違う。崩れの当たりをつけるのに使い、渡す前の目視は Keynote の PDF で行う。
色はテーマから、地色とレイアウトの飾り（タイトルの罫など）はレイアウトから読む。画像は描かない。画像にするには
consulting-pptx-skill の Playwright（npm run setup）が要る。
"""
import html
import subprocess
import sys
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

PX = 96
SCHEME = {}   # テーマの色。main() が scheme_of() で埋める


def scheme_of(prs):
    """テーマの色の名前と HEX の対応。bg1/tx1/bg2/tx2 は lt1/dk1/lt2/dk2 の別名。"""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from lxml import etree
    theme = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
    out = {}
    for el in theme.find(".//" + qn("a:clrScheme")):
        c = el[0]
        out[el.tag.split("}")[1]] = c.get("lastClr") or c.get("val")
    for a, b in (("bg1", "lt1"), ("tx1", "dk1"), ("bg2", "lt2"), ("tx2", "dk2")):
        out[a] = out[b]
    return out


def bg_of(slide):
    """スライドかレイアウトかマスターが持つ地色。無ければ白。"""
    for part in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        bg = part._element.cSld.find(qn("p:bg"))
        if bg is not None and bg.find(qn("p:bgPr")) is not None:
            return color_of(bg.find(qn("p:bgPr"))) or "#fff"
    return "#fff"


def emu(v):
    return v / 914400 * PX


def color_of(fill_parent):
    """<a:solidFill> を CSS の色に。"""
    if fill_parent is None:
        return None
    sf = fill_parent.find(qn("a:solidFill"))
    if sf is None:
        return None
    c = sf[0]
    base = c.get("val")
    rgb = base if c.tag == qn("a:srgbClr") else SCHEME.get(base, "888888")
    r, g, b = (int(rgb[i:i + 2], 16) for i in (0, 2, 4))
    a = 1.0
    mod = {el.tag.split("}")[1]: int(el.get("val")) for el in c}
    if "lumMod" in mod:
        f = mod["lumMod"] / 100000
        off = mod.get("lumOff", 0) / 100000 * 255
        r, g, b = (min(255, int(x * f + off)) for x in (r, g, b))
    if "alpha" in mod:
        a = mod["alpha"] / 100000
    return f"rgba({r},{g},{b},{a})"


def text_html(tx, w, h, x, y, rot=0):
    bp = tx.find(qn("a:bodyPr"))
    ins = [emu(int(bp.get(k, d))) for k, d in (("lIns", 91440), ("tIns", 45720), ("rIns", 91440), ("bIns", 45720))]
    anchor = {"t": "flex-start", "ctr": "center", "b": "flex-end"}[bp.get("anchor", "t")]
    paras = []
    for p in tx.findall(qn("a:p")):
        ppr = p.find(qn("a:pPr"))
        al = {"l": "left", "ctr": "center", "r": "right"}.get(ppr.get("algn", "l") if ppr is not None else "l", "left")
        mar = emu(int(ppr.get("marL", 0))) if ppr is not None else 0
        ind = emu(int(ppr.get("indent", 0))) if ppr is not None else 0
        bullet = ppr is not None and ppr.find(qn("a:buChar")) is not None
        sa = 0
        if ppr is not None and ppr.find(qn("a:spcAft")) is not None:
            sa = int(ppr.find(qn("a:spcAft"))[0].get("val")) / 100 / 72 * PX
        runs, first_sz, first_c = [], None, "#333"
        for el in p:
            if el.tag == qn("a:r"):
                rp = el.find(qn("a:rPr"))
                sz = int(rp.get("sz", 1000)) / 100 if rp is not None else 10
                first_sz = first_sz or sz
                col = color_of(rp) or "rgba(0,0,0,1)"
                first_c = first_c if runs else col
                bold = "700" if rp is not None and rp.get("b") == "1" else "400"
                runs.append(f'<span style="font-size:{sz / 72 * PX:.2f}px;font-weight:{bold};color:{col}">'
                            f'{html.escape(el.find(qn("a:t")).text or "")}</span>')
            elif el.tag == qn("a:br"):
                runs.append("<br>")
        if not runs:
            end = p.find(qn("a:endParaRPr"))
            sz = int(end.get("sz", 1000)) / 100 if end is not None else 10
            runs.append(f'<span style="font-size:{sz / 72 * PX:.2f}px">&nbsp;</span>')
        bl = (f'<span style="position:absolute;left:{ind:.1f}px;color:{first_c}">•</span>' if bullet else "")
        fs = (first_sz or 10) / 72 * PX
        paras.append(f'<div style="position:relative;font-size:{fs:.2f}px;text-align:{al};padding-left:{mar:.1f}px;margin-bottom:{sa:.1f}px">'
                     f'{bl}{"".join(runs)}</div>')
    tr = f"transform:rotate({rot}deg);" if rot else ""
    return (f'<div style="position:absolute;left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;{tr}'
            f'box-sizing:border-box;padding:{ins[1]:.1f}px {ins[2]:.1f}px {ins[3]:.1f}px {ins[0]:.1f}px;'
            f'display:flex;flex-direction:column;justify-content:{anchor};line-height:1.2">{"".join(paras)}</div>')


def shape_svg(sp, x, y, w, h):
    spPr = sp.find(qn("p:spPr"))
    geom = spPr.find(qn("a:prstGeom"))
    prst = geom.get("prst") if geom is not None else "rect"
    fill = color_of(spPr) or "none"
    xfrm = spPr.find(qn("a:xfrm"))
    rot = int(xfrm.get("rot", 0)) / 60000 if xfrm is not None else 0
    adj = 0.5
    gd = geom.find(qn("a:avLst")).find(qn("a:gd")) if geom is not None and geom.find(qn("a:avLst")) is not None else None
    if gd is not None:
        adj = int(gd.get("fmla").split()[1]) / 100000
    d = min(w, h) * adj
    pts = {
        "rect": [(0, 0), (w, 0), (w, h), (0, h)],
        "triangle": [(w / 2, 0), (w, h), (0, h)],
        "homePlate": [(0, 0), (w - d, 0), (w, h / 2), (w - d, h), (0, h)],
        "chevron": [(0, 0), (w - d, 0), (w, h / 2), (w - d, h), (0, h), (d, h / 2)],
    }.get(prst)
    tr = f'transform="rotate({rot} {w / 2} {h / 2})"' if rot else ""
    if prst == "ellipse":
        body = f'<ellipse cx="{w / 2}" cy="{h / 2}" rx="{w / 2}" ry="{h / 2}" fill="{fill}"/>'
    elif pts:
        body = f'<polygon points="{" ".join(f"{a:.1f},{b:.1f}" for a, b in pts)}" fill="{fill}" {tr}/>'
    else:
        body = f'<rect width="{w}" height="{h}" fill="none" stroke="red"/>'
    return (f'<svg style="position:absolute;left:{x:.1f}px;top:{y:.1f}px;overflow:visible" width="{w:.1f}" '
            f'height="{h:.1f}">{body}</svg>')


def line_svg(cx, x, y, w, h):
    spPr = cx.find(qn("p:spPr"))
    xfrm = spPr.find(qn("a:xfrm"))
    fx, fy = xfrm.get("flipH") == "1", xfrm.get("flipV") == "1"
    ln = spPr.find(qn("a:ln"))
    col = color_of(ln) or "#808080"
    lw = int(ln.get("w", 12700)) / 12700 / 72 * PX
    x1, x2 = (w, 0) if fx else (0, w)
    y1, y2 = (h, 0) if fy else (0, h)
    mk = ""
    if ln.find(qn("a:tailEnd")) is not None:
        mk = 'marker-end="url(#arr)"'
    return (f'<svg style="position:absolute;left:{x - 6:.1f}px;top:{y - 6:.1f}px;overflow:visible" width="{w + 12:.1f}" '
            f'height="{h + 12:.1f}"><defs><marker id="arr" markerWidth="6" markerHeight="6" refX="5" refY="3" '
            f'orient="auto"><path d="M0,0 L6,3 L0,6 z" fill="{col}"/></marker></defs>'
            f'<line x1="{x1 + 6}" y1="{y1 + 6}" x2="{x2 + 6}" y2="{y2 + 6}" stroke="{col}" stroke-width="{lw:.2f}" {mk}/></svg>')


def walk(shapes, out, idx):
    for sh in shapes:
        el = sh._element
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP:
            walk(sh.shapes, out, idx)
            continue
        x, y, w, h = emu(sh.left or 0), emu(sh.top or 0), emu(sh.width or 0), emu(sh.height or 0)
        if el.tag == qn("p:cxnSp"):
            out.append(line_svg(el, x, y, w, h))
            continue
        if el.tag != qn("p:sp"):
            continue
        if sh.is_placeholder and sh.placeholder_format.type == 13:
            out.append(f'<div style="position:absolute;left:{x:.1f}px;top:{y:.1f}px;width:{w:.1f}px;height:{h:.1f}px;'
                       f'display:flex;align-items:center;justify-content:flex-end;font-size:{8 / 72 * PX:.1f}px;'
                       f'color:#7f7f7f">{idx}</div>')
            continue
        out.append(shape_svg(el, x, y, w, h))
        tx = el.find(qn("p:txBody"))
        if tx is not None:
            rot = (sh.rotation or 0)
            out.append(text_html(tx, w, h, x, y, rot))


def main():
    src, outdir = sys.argv[1], Path(sys.argv[2])
    prs = Presentation(src)
    SCHEME.update(scheme_of(prs))
    W, H = emu(prs.slide_width), emu(prs.slide_height)
    pages = []
    for i, s in enumerate(prs.slides, 1):
        out = []
        walk([sh for sh in s.slide_layout.shapes if not sh.is_placeholder], out, i)
        walk(s.shapes, out, i)
        pages.append(f'<div class="pg" style="position:relative;width:{W}px;height:{H}px;background:{bg_of(s)};'
                     f'margin:0 0 12px 0;overflow:hidden">{"".join(out)}</div>')
    doc = ('<!doctype html><meta charset="utf-8"><body style="margin:0;background:#999;'
           'font-family:\'BIZ UDGothic\',\'Hiragino Sans\',sans-serif;color:#333">' + "".join(pages) + "</body>")
    outdir.mkdir(parents=True, exist_ok=True)
    page = outdir / "preview.html"
    page.write_text(doc, encoding="utf8")
    sys.path.insert(0, str(Path(__file__).resolve().parent))
    from house import CONSULTING
    full = outdir / "preview.png"
    subprocess.run(["node", str(Path(__file__).resolve().parent / "shot.mjs"), str(CONSULTING), str(page), str(full),
                    str(int(W))], check=True)
    from PIL import Image
    im = Image.open(full)
    ph = (int(H) + 12) * 2   # 1 ページの高さ（余白 12px、2 倍の解像度）
    n = round(im.height / ph)
    for k in range(0, n, 3):
        out = outdir / f"preview{k // 3 + 1}.png"
        im.crop((0, k * ph, im.width, min(im.height, (k + 3) * ph))).save(out)
        print(out, f"（{k + 1}〜{min(k + 3, n)} ページ）")


if __name__ == "__main__":
    main()
