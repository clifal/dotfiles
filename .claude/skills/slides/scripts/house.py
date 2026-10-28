#!/usr/bin/env python3
"""テンプレート（../assets/template.pptx）のレイアウトの上にページを組む部品。

デッキごとの build.py から import する（例: ../examples/sample/build.py）。
組み方と色の決まりは ../design.md、テンプレートと道具の癖は ../quirks.md にある。
"""
import copy
import datetime
import json
import os
import sys
import unicodedata
from pathlib import Path

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE.parent / "assets" / "template.pptx"
SKIN = HERE / "house.skin.json"
CONSULTING = Path(os.environ.get("CONSULTING_PPTX_SKILL", Path.home() / ".claude" / "skills" / "consulting-pptx-skill"))
if not (CONSULTING / "scripts" / "deck_pptx.py").exists():
    sys.exit(f"consulting-pptx-skill が見つからない: {CONSULTING}（SKILL.md の「準備」を見る）")
sys.path.insert(0, str(CONSULTING / "scripts"))

from deck_pptx import Deck, P, E, sub, OVAL  # noqa: E402,F401
from lxml import etree  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402,F401
from pptx.enum.shapes import PP_PLACEHOLDER as PH  # noqa: E402
from pptx.oxml.ns import qn  # noqa: E402

TXT = "tx1"                    # 本文の文字（テーマの dk1 = 333333）
WHITE = "FFFFFF"
SUB = "bg1|lumMod=50000"       # フッターとページ番号の灰
LIGHT = "bg1|lumMod=85000"     # 目盛りと枠の薄い灰
GRAY = "bg1|lumMod=75000"      # グラフで目を向けない区分
PALE = "bg1|lumMod=95000"      # 色を付けない箱の地
# 区分の 3 色（カラーユニバーサルデザイン推奨配色）。_D は白い文字を載せる濃い色、_T は面に敷く薄い色
BLUE, BLUE_D, BLUE_T = "accent1", "accent1|lumMod=75000", "accent1|lumMod=20000|lumOff=80000"
GREEN, GREEN_D, GREEN_T = "accent2", "accent2|lumMod=60000", "accent2|lumMod=20000|lumOff=80000"
ORANGE, ORANGE_T = "accent3", "accent3|lumMod=20000|lumOff=80000"
SKY = "accent4"                # 章扉の番号
X0, X1, TOP, BOTTOM = 0.4, 9.6, 0.85, 5.0   # 本文の範囲（インチ）
SRC_Y = 5.06                                # 出典行の上端（フッターの 1 行上）
SIZES = {"表紙": 26, "章扉": 28, "本文": 16}  # 各レイアウトのタイトルの大きさ（make_template.py と同じ）


class HouseDeck(Deck):
    """文字色の既定をテーマの本文の色にする。"""

    def _rpr(self, parent, sz, b, c):
        super()._rpr(parent, sz, b, c or TXT)


def _layout(prs, name):
    for m in prs.slide_masters:
        for lay in m.slide_layouts:
            if lay.name == name:
                return lay
    raise KeyError(name)


class Builder:
    """表紙・章扉・本文のページを、テンプレートのレイアウトから順に足していく。

    footer はページの左下に出す 1 行（例「© 2026 〇〇株式会社」）。空なら出さない。"""

    def __init__(self, footer="", template=TEMPLATE, skin=SKIN):
        if not isinstance(skin, dict):
            skin = json.loads(Path(skin).read_text(encoding="utf8"))
        self.d = HouseDeck(str(template), skin)   # テンプレートの見本のページは外れる
        self.lay = {n: _layout(self.d.prs, n) for n in SIZES}
        num = next(p for p in self.lay["本文"].placeholders if p.placeholder_format.type == PH.SLIDE_NUMBER)
        self.num_sp = copy.deepcopy(num._element)
        self.footer = footer
        self.pages = {}

    # ---------------------------------------------------------------- 枠
    def _fill(self, ph, lines, sz, b=False, c=None, factor=None):
        """プレースホルダーに文字を入れる。書体・大きさ・色を run に書く（Keynote と preview.py の見た目をそろえる）。
        factor を渡すと、行の間隔を文字の大きさの factor 倍に固定する。"""
        txBody = ph.text_frame._txBody
        for p in txBody.findall(qn("a:p")):
            txBody.remove(p)
        for ln in lines:
            p = sub(txBody, "a:p")
            if factor:
                sub(sub(sub(p, "a:pPr"), "a:lnSpc"), "a:spcPts", val=str(int(sz * factor * 100)))
            r = sub(p, "a:r")
            self.d._rpr(r, sz, b, c)
            sub(r, "a:t").text = ln

    def cover(self, org, title, by="", date=None):
        """表紙。最初に呼ぶ。宛名（org）は題の上、日付と作成者（by）は題の下に置く。日付を省くと今日。
        題は cover_lines で行に分ける。"""
        if len(self.d.prs.slides):
            sys.exit("cover() は最初のページとして呼ぶ")
        s = self.d.prs.slides.add_slide(self.lay["表紙"])
        date = date or datetime.date.today().strftime("%Y/%m/%d")
        for ph in s.placeholders:
            if ph.placeholder_format.type == PH.CENTER_TITLE:
                self._fill(ph, cover_lines(title), SIZES["表紙"], b=True, factor=1.3)
            else:
                self._fill(ph, [date, by] if by else [date], 12, c="bg1|lumMod=35000")
        if org:
            self.d.text(s, 0.7, 0.95, 7.5, 0.4, P(org, sz=14), anchor="b", name="宛名")
        return s

    def divider(self, no, title, lead=""):
        """章扉。地色と飾りはレイアウトが持つ。番号は 2 桁で題の上に置く。ページ番号は出さないが、ページには数える。"""
        s = self.d.prs.slides.add_slide(self.lay["章扉"])
        for ph in list(s.placeholders):
            if ph.placeholder_format.type == PH.TITLE:
                self._fill(ph, [title], SIZES["章扉"], b=True, c=WHITE)
            elif lead:
                self._fill(ph, [lead], 12, c=WHITE)
            else:
                ph._element.getparent().remove(ph._element)
        self.d.text(s, 0.7, 0.95, 3.0, 0.95, P(f"{no:02d}", sz=44, b=True, c=SKY), anchor="b", name="章番号")
        return s

    def plan(self, keys):
        """ページの並び（表紙から順）を先に決める。ref() はこの並びからページ番号を数える。"""
        self.pages = {k: i for i, k in enumerate(keys, 1)}

    def ref(self, *keys):
        """ページの参照「P.n」。続く番号は「〜」で、離れた番号は「、」でつなぐ。"""
        runs = []
        for n in sorted(self.pages[k] for k in keys):
            if runs and n == runs[-1][1] + 1:
                runs[-1][1] = n
            else:
                runs.append([n, n])
        return "P." + "、".join(f"{a}〜{z}" if z > a else str(a) for a, z in runs)

    def slide(self, title):
        """本文のページ。タイトルの罫と印はレイアウトが持つ。ページ番号の枠はレイアウトから写す。"""
        s = self.d.prs.slides.add_slide(self.lay["本文"])
        self._fill(s.shapes.title, [title], SIZES["本文"], b=True)
        el = copy.deepcopy(self.num_sp)
        el.find(".//" + qn("p:cNvPr")).set("id", str(s.shapes._next_shape_id))
        s.shapes._spTree.append(el)
        if self.footer:
            self.d.text(s, X0, 5.3, 6.0, 0.22, P(self.footer, sz=7, c=SUB), anchor="m", name="フッター")
        return s

    def save(self, out, title):
        """作成者と最終更新者を空にして保存する。日本語のファイル名は NFC にそろえる（macOS は濁点を分けて保存し、Windows で化ける）。"""
        n = len(self.d.prs.slides)
        if self.pages and n != len(self.pages):
            sys.exit(f"plan() の並びは {len(self.pages)} ページで、組んだページは {n} 枚（参照の P.n がずれる）")
        cp = self.d.prs.core_properties
        cp.author, cp.last_modified_by, cp.title = "", "", title
        out = Path(unicodedata.normalize("NFC", str(out)))
        out.parent.mkdir(parents=True, exist_ok=True)
        self.d.save(out)
        return out

    # ---------------------------------------------------------------- 部品
    def source(self, s, text, lines=1):
        """出典・注記の行（フッターの 1 行上 — slide-rules §4.47）。"""
        h = 0.2 + 0.13 * (lines - 1)
        self.d.text(s, X0, SRC_Y + 0.2 - h, X1 - X0, h, P(text, sz=7.5), anchor="b")

    def head(self, s, x, y, w, label):
        """節の見出しと下罫。戻り値は下罫の下の y。"""
        self.d.head(s, x, y, w, label, sz=11, h=0.3)
        self.d.rule(s, y + 0.33, x0=x, x1=x + w, heavy=True)
        return y + 0.4

    def circle(self, s, x, y, n, dia=0.26):
        return self.d.tbox(s, x, y, dia, dia, P(str(n), sz=9, b=True, c=WHITE, al="c"), fill=BLUE_D, kind=OVAL,
                           pad=(0, 0, 0, 0), name=f"番号 {n}")

    def card(self, s, x, y, w, h, name, desc, fill=WHITE, c=None, stroke=None):
        """名前と説明を中央に置いた箱（体制図など）。stroke を渡すと、塗らずに枠だけの箱にする（§5.3）。"""
        box = self.d.tbox(s, x, y, w, h, [P(name, sz=9.5, b=True, c=c, al="c", sa=2), P(desc, sz=8.5, c=c, al="c")],
                          fill=fill, pad=(0.1, 0.05, 0.1, 0.05))
        if stroke:
            self.d._paint(box, None, stroke, 0.75)
        return box

    def legend(self, s, right, y, items):
        """色の凡例を、right を右端にして y の高さへ右揃えで置く（slide-rules §5.18）。items は (色, 名前)。"""
        x = right
        for fill, label in reversed(items):
            tw = len(label) * 0.118 + 0.08
            x -= tw
            self.d.text(s, x, y, tw, 0.2, P(label, sz=8.5), anchor="m")
            x -= 0.17
            self.d.box(s, x, y + 0.04, 0.12, 0.12, fill=fill)
            x -= 0.2


def deck_path(folder, title):
    """デッキのファイル名は題と同じにする（NFC。パスに使えない / と : は全角に替える）。"""
    name = unicodedata.normalize("NFC", title).replace("/", "／").replace(":", "：")
    return Path(folder) / f"{name}.pptx"


def cover_lines(title, width=20):
    """表紙の題を、半角の空白の位置で、1 行が全角 width 字ほどに収まるように分ける（右の面にかからない幅）。"""
    def w(t):
        return sum(1 if unicodedata.east_asian_width(ch) in "WFA" else 0.5 for ch in t)
    lines = []
    for word in title.split(" "):
        if lines and w(lines[-1] + " " + word) <= width:
            lines[-1] += " " + word
        else:
            lines.append(word)
    return lines


def pitch(shape, factor=1.25):
    """段落ごとに行の間隔を文字の大きさの factor 倍に固定する（Keynote は和文の行を高めに描く）。"""
    for p in shape.text_frame.paragraphs:
        r = p._p.find(qn("a:r"))
        sz = int(r.find(qn("a:rPr")).get("sz")) if r is not None else 1000
        ln = etree.Element(qn("a:lnSpc"))
        etree.SubElement(ln, qn("a:spcPts")).set("val", str(int(sz * factor)))
        p._p.find(qn("a:pPr")).insert(0, ln)
    return shape


def bl(items, sz=9.5):
    """箇条書きの段落の並び。"""
    return [P(t, sz=sz, bullet=True) for t in items]
