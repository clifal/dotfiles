#!/usr/bin/env python3
"""PPTX の文言を、ページごと・行（同じ高さの箱）ごとにテキストへ書き出す。yomiyasu で推敲するときに使う。

  python3 -I ~/.claude/skills/slides/scripts/export_text.py <pptx> <出力.md> [ページ番号,...]

タイトルは「タイトル　」を頭に付けた行にする（「#」を付けると、yomiyasu が見出しとして短く直そうとする）。
箱の中の改行は行を分けて残す。番号の丸、章扉の番号、フッター、ページ番号は書き出さない。
"""
import sys

from pptx import Presentation
from pptx.oxml.ns import qn


def lines_of(p):
    buf, cur = [], ""
    for el in p._p:
        if el.tag == qn("a:r"):
            cur += el.find(qn("a:t")).text or ""
        elif el.tag == qn("a:br"):
            buf.append(cur)
            cur = ""
    buf.append(cur)
    return [b.strip() for b in buf if b.strip()]


def main():
    src, dst = sys.argv[1], sys.argv[2]
    only = {int(x) for x in sys.argv[3].split(",")} if len(sys.argv) > 3 else None
    prs = Presentation(src)
    out = [f"{prs.core_properties.title or src}（スライドの文言）", ""]
    for i, s in enumerate(prs.slides, 1):
        if only and i not in only:
            continue
        title = s.shapes.title
        tid = title.shape_id if title is not None else None
        out += ["---", "", f"P.{i}", ""]
        if tid is not None and i > 1:
            out += ["タイトル　" + title.text_frame.text.replace("\n", ""), ""]
        rest = [sh for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip() and sh.shape_id != tid
                and sh.name not in ("フッター", "章番号") and not sh.name.startswith("番号")
                and not (sh.is_placeholder and sh.placeholder_format.type == 13)]
        rest.sort(key=lambda sh: (sh.top, sh.left))
        rows, cur, ctop = [], [], None
        for sh in rest:   # 上端が 0.12in 以内の箱を同じ行とみなす
            if ctop is None or sh.top - ctop > 914400 * 0.12:
                if cur:
                    rows.append(cur)
                cur, ctop = [], sh.top
            cur.append(sh)
        if cur:
            rows.append(cur)
        for row in rows:
            for sh in sorted(row, key=lambda x: x.left):
                for p in sh.text_frame.paragraphs:
                    ls = lines_of(p)
                    if not ls:
                        continue
                    ppr = p._p.find(qn("a:pPr"))
                    bullet = ppr is not None and ppr.find(qn("a:buChar")) is not None
                    out.append(("- " if bullet else "") + ls[0])
                    out += ls[1:]
            out.append("")
    open(dst, "w", encoding="utf8").write("\n".join(out))
    print("wrote", dst)


if __name__ == "__main__":
    main()
